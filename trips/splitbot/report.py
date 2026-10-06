#!/usr/bin/env python3
"""Trip reports: the nightly breakdown, charts (PNG) and the Excel log.

All figures come from splitlog.py (SGD cents), so the report, the chat replies
and the workbook can never disagree. Files go to ~/.splitbot/out/ (outside the
repo); expenses.csv stays the record and the workbook is regenerated from it.

Usage (from the repo root):
  report.py daily --chat ID [--send] [--skip-empty] [--no-mark]
      Nightly breakdown: text + chart + Excel. --send posts all three to the
      group and records the time, so the next report lists only what's new.
  report.py nightly [--at 21:45]
      For every open trip (not in planning) whose local time is past --at and
      that has had no nightly report today: post the breakdown (if anything new was logged)
      and mark the day done. Run from the hourly watchdog; idempotent.
  report.py chart --chat ID [--send]      the chart image only
  report.py excel --chat ID [--send]      the Excel workbook only
  report.py endtrip --chat ID --send      /endtrip in one step: close the trip, then
      post the final settle-up, chart, Excel log and Wrapped PDF to the group.
  report.py catalogue --chat ID [--send]  the plan as a PDF (cover, at a glance, day by
      day itinerary with Maps links, the shortlist by kind, bookings, decisions).
  report.py wrapped --chat ID [--send]    the trip PDF: Wrapped-style cards, then
      the full expense report. Works on a closed trip too (the group's latest).
  Add --to CHAT to deliver somewhere other than the trip's group, e.g. the
  owner's private chat with the bot (--chat still picks the trip).
"""

from __future__ import annotations

import argparse
import html
import math
import os
import subprocess
import sys
from collections import defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import splitlog as sl  # noqa: E402

# Category -> fixed palette slot (colour follows the category, never its rank).
# Reference palette, light mode; validated as a set (adjacent CVD dE >= 9.1).
CAT_COLOUR = {"food": "#2a78d6", "drinks": "#4a3aa7", "transport": "#eb6834", "lodging": "#1baf7a",
              "activities": "#eda100", "shopping": "#e87ba4", "other": "#008300"}
POS, NEG = "#2a78d6", "#e34948"            # diverging pair: is owed / owes
INK, INK2, MUTED, GRID, AXIS, SURFACE = "#0b0b0b", "#52514e", "#898781", "#e1e0d9", "#c3c2b7", "#fcfcfb"


def out_dir(trip: dict) -> Path:
    d = Path(os.environ.get("SPLITBOT_HOME", Path.home() / ".splitbot")) / "out" / trip["id"]
    d.mkdir(parents=True, exist_ok=True)
    return d


def cents(r: dict) -> int:
    return round(float(r["amount_sgd"]) * 100)


def expenses(trip: dict) -> list[dict]:
    return [r for r in sl.trip_rows(trip) if r["kind"] == "expense"]


def rate_note(trip: dict) -> str:
    """How this trip's payments were converted, named from the rates actually used."""
    srcs = {r["fx_source"] for r in sl.trip_rows(trip)} - {"identity", ""}
    parts = []
    if any(s.startswith("ECB") for s in srcs):
        parts.append("ECB rates on the payment day")
    pairs = sorted({s.split()[1].removesuffix("=X") for s in srcs if s.startswith("Yahoo ")})
    if pairs:
        parts.append("Yahoo Finance " + ", ".join(f"{p[:3]}/{p[3:]}" for p in pairs)
                     + " rates when logged")
    if "given" in srcs:
        parts.append("the card's own rate where someone gave it")
    return "converted to SGD at " + "; ".join(parts) if parts else "all amounts in SGD"


def day_label(d: str) -> str:
    x = date.fromisoformat(d)
    return f"{x:%a} {x.day} {x:%b}"


def trip_days(trip: dict, rows: list[dict]) -> list[str]:
    """Every day from the first payment (or trip start) to the last or today."""
    today = trip.get("closed") if trip["status"] == "closed" and trip.get("closed") else sl.trip_today(trip)
    dates = [r["date"] for r in rows] + [today]
    start = min(dates + [trip["created"]]) if rows else today
    end = max(dates)
    d0, d1 = date.fromisoformat(start), date.fromisoformat(end)
    return [(d0 + timedelta(days=i)).isoformat() for i in range((d1 - d0).days + 1)]


def by_day_cat(rows: list[dict]) -> dict[str, dict[str, int]]:
    out: dict[str, dict[str, int]] = defaultdict(lambda: defaultdict(int))
    for r in rows:
        out[r["date"]][r.get("category") or "other"] += cents(r)
    return out


def local_total(rows: list[dict]) -> str:
    per = defaultdict(float)
    for r in rows:
        per[r["currency"]] += float(r["amount"])
    return " + ".join(sl.money(v, c) for c, v in sorted(per.items(), key=lambda kv: -kv[1]))


# ------------------------------------------------------------------------ text
def daily_text(trip: dict, day: str | None = None) -> tuple[str, int]:
    """The nightly message and how many entries are new since the last report."""
    day = day or sl.trip_today(trip)
    rows = expenses(trip)
    live = sl.trip_rows(trip)
    since = trip.get("last_report")
    new = [r for r in live if (r["logged_at"] > since if since else r["date"] == day)]
    today = [r for r in rows if r["date"] == day]
    earlier = [r for r in new if r["date"] != day and r["kind"] == "expense"]
    total = sum(cents(r) for r in today)

    lines = [f"{trip['name']} · {day_label(day)}"]
    if today:
        lines.append(f"Today: {sl.sgd(total)} across {len(today)} payment{'s' * (len(today) != 1)}"
                     f" ({local_total(today)})")
        cat = defaultdict(int)
        for r in today:
            cat[r.get("category") or "other"] += cents(r)
        lines.append(" · ".join(f"{c.capitalize()} {sl.sgd(v)}"
                                for c, v in sorted(cat.items(), key=lambda kv: -kv[1])))
        paid = defaultdict(int)
        for r in today:
            paid[r["payer"]] += cents(r)
        lines.append("Paid today: " + ", ".join(f"{k} {sl.sgd(v)}"
                                                for k, v in sorted(paid.items(), key=lambda kv: -kv[1])))
        share = defaultdict(int)
        for r in today:
            for k, v in sl.read_shares(r["shares_sgd"]).items():
                share[k] += v
        lines.append("Each person's share today: " + ", ".join(
            f"{k} {sl.sgd(v)}" for k, v in sorted(share.items(), key=lambda kv: -kv[1])))
    else:
        lines.append("No payments logged today.")
    if earlier:
        lines.append("Also logged for earlier days: " + "; ".join(
            f"#{r['entry']} {r['desc']} ({day_label(r['date'])}) {sl.sgd(cents(r))}" for r in earlier))

    days = trip_days(trip, rows)
    spent = sum(cents(r) for r in rows)
    lines += ["", f"Trip so far: {sl.sgd(spent)} over {len(days)} day{'s' * (len(days) != 1)}"
                  f" (about {sl.sgd(round(spent / max(len(days), 1)))} a day)"]
    nets = {k: v["net"] for k, v in sl.nets(trip).items()}
    lines.append("Balances: " + " · ".join(
        f"{k} {'+' if v > 0 else '−' if v < 0 else '±'}{sl.sgd(abs(v))}"
        for k, v in sorted(nets.items(), key=lambda kv: -kv[1])))
    plan = sl.settle_plan(nets)
    lines.append("To settle now: " + (" · ".join(f"{d} → {c} {sl.sgd(x)}" for d, c, x in plan)
                                      if plan else "everyone is square"))
    lines.append(f"(+ is owed, − owes; {rate_note(trip)})")
    return "\n".join(lines), len(new)


# ----------------------------------------------------------------------- chart
def nice_step(top: float) -> float:
    raw = top / 4 if top > 0 else 1
    mag = 10 ** math.floor(math.log10(raw))
    return next(m * mag for m in (1, 2, 2.5, 5, 10) if m * mag >= raw)


def esc(s: str) -> str:
    return html.escape(str(s), quote=True)


def svg_daily(trip: dict, rows: list[dict], highlight: str) -> str:
    days = trip_days(trip, rows)
    data = by_day_cat(rows)
    cats = [c for c in CAT_COLOUR if any(data[d].get(c) for d in days)] or ["other"]
    totals = [sum(data[d].values()) / 100 for d in days]
    W, H, L, R, T, B = 820, 300, 64, 16, 28, 44
    step = nice_step(max(totals + [1]))
    top = step * max(1, math.ceil(max(totals + [0.01]) / step))
    pw = W - L - R
    slot = pw / len(days)
    bw = min(44, slot * 0.62)
    y = lambda v: T + (H - T - B) * (1 - v / top)  # noqa: E731
    o = [f"<svg viewBox='0 0 {W} {H}' width='{W}' height='{H}' xmlns='http://www.w3.org/2000/svg'>"]
    v = 0.0
    while v <= top + 1e-9:
        o.append(f"<line x1='{L}' x2='{W - R}' y1='{y(v):.1f}' y2='{y(v):.1f}' stroke='{GRID if v else AXIS}'/>")
        o.append(f"<text x='{L - 8}' y='{y(v) + 4:.1f}' text-anchor='end' fill='{MUTED}' font-size='12'>"
                 f"{v:,.0f}</text>")
        v += step
    show_every = max(1, math.ceil(len(days) / 14))
    for i, d in enumerate(days):
        cx = L + slot * (i + 0.5)
        x0 = cx - bw / 2
        base = 0.0
        segs = [(c, data[d].get(c, 0) / 100) for c in cats if data[d].get(c)]
        for j, (c, val) in enumerate(segs):
            y_top, y_bot = y(base + val), y(base)
            gap = 2 if j else 0          # 2px surface gap between stacked segments
            h = max(y_bot - y_top - gap, 0.5)
            if j == len(segs) - 1:       # 4px rounded data-end on the top segment
                r = min(4, h)
                o.append(f"<path d='M{x0:.1f},{y_top + h:.1f} V{y_top + r:.1f} Q{x0:.1f},{y_top:.1f} "
                         f"{x0 + r:.1f},{y_top:.1f} H{x0 + bw - r:.1f} Q{x0 + bw:.1f},{y_top:.1f} "
                         f"{x0 + bw:.1f},{y_top + r:.1f} V{y_top + h:.1f} Z' fill='{CAT_COLOUR[c]}'/>")
            else:
                o.append(f"<rect x='{x0:.1f}' y='{y_top:.1f}' width='{bw:.1f}' height='{h:.1f}' "
                         f"fill='{CAT_COLOUR[c]}'/>")
            base += val
        if totals[i]:
            o.append(f"<text x='{cx:.1f}' y='{y(totals[i]) - 6:.1f}' text-anchor='middle' fill='{INK2}' "
                     f"font-size='12'>{totals[i]:,.0f}</text>")
        if i % show_every == 0 or d == highlight:
            wt = "700" if d == highlight else "400"
            fill = INK if d == highlight else MUTED
            o.append(f"<text x='{cx:.1f}' y='{H - B + 18}' text-anchor='middle' fill='{fill}' "
                     f"font-size='12' font-weight='{wt}'>{esc(day_label(d)[:-4])}</text>")
    o.append("</svg>")
    legend = "".join(f"<span class='key'><i style='background:{CAT_COLOUR[c]}'></i>{c.capitalize()}</span>"
                     for c in cats)
    return f"<div class='legend'>{legend}</div>" + "".join(o)


def svg_balances(trip: dict) -> str:
    nets = sorted(((k, v["net"] / 100) for k, v in sl.nets(trip).items()), key=lambda kv: -kv[1])
    W, row, L = 820, 34, 120
    H = row * len(nets) + 16
    mx = max([abs(v) for _, v in nets] + [1])
    # Reserve room for the longest value label on each side, so a long bar's
    # label never runs into the names column or off the right edge.
    tags = [f"is owed {sl.sgd(round(v * 100))}" if v > 0 else f"owes {sl.sgd(round(-v * 100))}"
            for _, v in nets if v]
    lab = max([len(t) for t in tags] + [6]) * 7.2 + 12
    half = (W - L - 2 * lab) / 2
    mid = L + lab + half
    o = [f"<svg viewBox='0 0 {W} {H}' width='{W}' height='{H}' xmlns='http://www.w3.org/2000/svg'>",
         f"<line x1='{mid}' x2='{mid}' y1='0' y2='{H}' stroke='{AXIS}'/>"]
    for i, (name, v) in enumerate(nets):
        yc = 8 + row * i + row / 2
        w = abs(v) / mx * half
        colour = POS if v > 0 else NEG
        x = mid if v >= 0 else mid - w
        if w > 0.5:
            o.append(f"<rect x='{x:.1f}' y='{yc - 9:.1f}' width='{w:.1f}' height='18' rx='4' fill='{colour}'/>")
        o.append(f"<text x='{L - 10}' y='{yc + 4:.1f}' text-anchor='end' fill='{INK}' font-size='13'>{esc(name)}</text>")
        tag = f"is owed {sl.sgd(round(v * 100))}" if v > 0 else \
            f"owes {sl.sgd(round(-v * 100))}" if v < 0 else "square"
        tx, anchor = (mid + w + 8, "start") if v >= 0 else (mid - w - 8, "end")
        o.append(f"<text x='{tx:.1f}' y='{yc + 4:.1f}' text-anchor='{anchor}' fill='{INK2}' "
                 f"font-size='12'>{tag}</text>")
    o.append("</svg>")
    return "".join(o)


def chart_html(trip: dict, day: str) -> str:
    rows = expenses(trip)
    spent = sum(cents(r) for r in rows)
    n_days = len(trip_days(trip, rows))
    return f"""<!doctype html><html><head><meta charset='utf-8'><style>
body {{ margin:0; background:{SURFACE}; font-family:'Inter','Noto Sans','DejaVu Sans',system-ui,sans-serif; }}
#viz {{ width:860px; padding:20px 20px 16px; background:{SURFACE}; color:{INK}; box-sizing:border-box; }}
h1 {{ font-size:20px; margin:0 0 2px; }} h2 {{ font-size:15px; margin:18px 0 6px; color:{INK}; }}
.sub {{ color:{INK2}; font-size:13px; margin:0 0 8px; }}
.legend {{ display:flex; gap:14px; flex-wrap:wrap; font-size:12px; color:{INK2}; margin:2px 0 4px; }}
.key i {{ display:inline-block; width:10px; height:10px; border-radius:2px; margin-right:5px; vertical-align:-1px; }}
.foot {{ color:{MUTED}; font-size:11px; margin-top:8px; }}
</style></head><body><div id='viz'>
<h1>{esc(trip['name'])}</h1>
<p class='sub'>{sl.sgd(spent)} spent over {n_days} day{'s' * (n_days != 1)} · as of {esc(day_label(day))}</p>
<h2>Spending per day (SGD), by category</h2>{svg_daily(trip, rows, day)}
<h2>Who is owed and who owes (SGD)</h2>{svg_balances(trip)}
<p class='foot'>{esc(rate_note(trip)[0].upper() + rate_note(trip)[1:])}. Settle-ups excluded from spending.</p>
</div></body></html>"""


def render_chart(trip: dict, day: str) -> Path:
    d = out_dir(trip)
    src, png = d / "chart.html", d / f"{trip['id']}-{day}.png"
    src.write_text(chart_html(trip, day), encoding="utf-8")
    env = dict(os.environ)
    if "NODE_PATH" not in env:
        env["NODE_PATH"] = subprocess.run(["npm", "root", "-g"], capture_output=True, text=True).stdout.strip()
    subprocess.run(["node", str(HERE / "render_png.js"), str(src), str(png), "860"],
                   check=True, capture_output=True, text=True, env=env)
    return png


# ----------------------------------------------------------------------- excel
def build_excel(trip: dict) -> Path:
    try:
        import openpyxl  # noqa: F401
    except ImportError:  # containers are rebuilt without it; install on first use
        subprocess.run([sys.executable, "-m", "pip", "install", "-q", "openpyxl"],
                       check=True, capture_output=True, text=True)
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="2A78D6")
    bold = Font(bold=True)
    members = [m["name"] for m in trip["members"]]
    live = sorted(sl.trip_rows(trip), key=lambda r: int(r["entry"]))
    wb = Workbook()

    def header(ws, cols, widths):
        ws.append(cols)
        for i, w in enumerate(widths, 1):
            c = ws.cell(row=1, column=i)
            c.font, c.fill = head, fill
            c.alignment = Alignment(vertical="center", wrap_text=True)
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"

    # Transactions: one row per payment, a column per person's share.
    ws = wb.active
    ws.title = "Transactions"
    cols = ["#", "Date", "Description", "Category", "Paid by", "Amount", "Currency",
            "SGD per unit", "SGD", "Split"] + [f"{m} share (SGD)" for m in members] + ["Logged by", "Rate source"]
    header(ws, cols, [5, 11, 28, 12, 12, 13, 9, 12, 11, 22] + [13] * len(members) + [13, 22])
    for r in live:
        shares = sl.read_shares(r["shares_sgd"])
        ws.append([int(r["entry"]), date.fromisoformat(r["date"]), r["desc"],
                   (r.get("category") or "other").capitalize(), r["payer"], float(r["amount"]),
                   r["currency"], float(r["fx_to_sgd"]), float(r["amount_sgd"]),
                   r["split"] if r["kind"] == "expense" else f"settle-up to {r['split']}"]
                  + [shares.get(m, 0) / 100 for m in members] + [r["logged_by"], r["fx_source"]])
        n = ws.max_row
        ws.cell(n, 2).number_format = "ddd d mmm yyyy"
        ws.cell(n, 6).number_format = "#,##0" if sl.dp(r["currency"]) == 0 else "#,##0.00"
        ws.cell(n, 8).number_format = "0.000000"
        for c in range(9, 11 + len(members)):
            ws.cell(n, c).number_format = "#,##0.00"
    last = ws.max_row
    if live:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{last}"
        ws.append([])
        tot = ["Total (spending, excl. settle-ups)", None, None, None, None, None, None, None,
               f'=SUMIFS(I2:I{last},D2:D{last},"<>Settle-up")', None]
        tot += [None] * len(members)
        ws.append(tot)
        ws.cell(ws.max_row, 1).font = bold
        ws.cell(ws.max_row, 9).font = bold
        ws.cell(ws.max_row, 9).number_format = "#,##0.00"

    # Daily: date x category, and each person's share of the day.
    exp = expenses(trip)
    days = trip_days(trip, exp)
    data = by_day_cat(exp)
    cats = [c for c in CAT_COLOUR if any(data[d].get(c) for d in days)] or ["other"]
    share_day: dict[str, dict[str, int]] = defaultdict(lambda: defaultdict(int))
    for r in exp:
        for k, v in sl.read_shares(r["shares_sgd"]).items():
            share_day[r["date"]][k] += v
    wd = wb.create_sheet("Daily")
    header(wd, ["Date"] + [c.capitalize() for c in cats] + ["Total"] + [f"{m}'s share" for m in members],
           [14] + [12] * (len(cats) + 1 + len(members)))
    for d in days:
        i = wd.max_row + 1
        wd.append([date.fromisoformat(d)] + [data[d].get(c, 0) / 100 for c in cats]
                  + [f"=SUM(B{i}:{get_column_letter(1 + len(cats))}{i})"]
                  + [share_day[d].get(m, 0) / 100 for m in members])
        wd.cell(i, 1).number_format = "ddd d mmm"
        for c in range(2, 3 + len(cats) + len(members)):
            wd.cell(i, c).number_format = "#,##0.00"
    n = wd.max_row
    wd.append(["Trip total"] + [f"=SUM({get_column_letter(c)}2:{get_column_letter(c)}{n})"
                                for c in range(2, 3 + len(cats) + len(members))])
    for c in range(1, 3 + len(cats) + len(members)):
        wd.cell(wd.max_row, c).font = bold
        wd.cell(wd.max_row, c).number_format = "#,##0.00"
    ch = BarChart()
    ch.type, ch.grouping, ch.overlap = "col", "stacked", 100
    ch.title, ch.y_axis.title = "Spending per day (SGD)", "SGD"
    ch.add_data(Reference(wd, min_col=2, max_col=1 + len(cats), min_row=1, max_row=n), titles_from_data=True)
    ch.set_categories(Reference(wd, min_col=1, min_row=2, max_row=n))
    for s, c in zip(ch.series, cats):
        s.graphicalProperties.solidFill = CAT_COLOUR[c][1:]
        s.graphicalProperties.line.noFill = True
    ch.height, ch.width = 9, 22
    wd.add_chart(ch, f"A{n + 3}")

    # Balances and settle-up.
    wbal = wb.create_sheet("Balances")
    header(wbal, ["Person", "Paid (SGD)", "Share (SGD)", "Net (SGD)", "Status"], [14, 13, 13, 13, 18])
    nets = sl.nets(trip)
    for name, v in sorted(nets.items(), key=lambda kv: -kv[1]["net"]):
        wbal.append([name, v["paid"] / 100, v["share"] / 100, v["net"] / 100,
                     "is owed" if v["net"] > 0 else "owes" if v["net"] < 0 else "square"])
        for c in (2, 3, 4):
            wbal.cell(wbal.max_row, c).number_format = "#,##0.00"
    nb = wbal.max_row
    bc = BarChart()
    bc.type, bc.title, bc.legend = "bar", "Net balance (SGD): + is owed, − owes", None
    bc.add_data(Reference(wbal, min_col=4, min_row=1, max_row=nb), titles_from_data=True)
    bc.set_categories(Reference(wbal, min_col=1, min_row=2, max_row=nb))
    bc.series[0].graphicalProperties.solidFill = POS[1:]
    bc.series[0].invertIfNegative = True
    bc.height, bc.width = 7, 16
    wbal.add_chart(bc, "G2")
    wbal.append([])
    wbal.append(["To settle up"])
    wbal.cell(wbal.max_row, 1).font = bold
    wbal.append(["From", "To", "SGD"])
    for c in (1, 2, 3):
        wbal.cell(wbal.max_row, c).font = bold
    for d_, c_, x in sl.settle_plan({k: v["net"] for k, v in nets.items()}):
        wbal.append([d_, c_, x / 100])
        wbal.cell(wbal.max_row, 3).number_format = "#,##0.00"

    # By category.
    wc = wb.create_sheet("By category")
    header(wc, ["Category", "SGD", "Share of spending", "Payments"], [14, 12, 16, 10])
    cat_tot = defaultdict(int)
    cat_n = defaultdict(int)
    for r in exp:
        cat_tot[r.get("category") or "other"] += cents(r)
        cat_n[r.get("category") or "other"] += 1
    total = sum(cat_tot.values()) or 1
    for c in sorted(cat_tot, key=lambda c: -cat_tot[c]):
        wc.append([c.capitalize(), cat_tot[c] / 100, cat_tot[c] / total, cat_n[c]])
        wc.cell(wc.max_row, 2).number_format = "#,##0.00"
        wc.cell(wc.max_row, 3).number_format = "0%"

    wi = wb.create_sheet("About")
    for row in [["Trip", trip["name"]], ["Local currency", trip["currency"]],
                ["Timezone", str(sl.trip_tz(trip))], ["Members", ", ".join(members)],
                ["Generated", datetime.now(sl.trip_tz(trip)).strftime("%Y-%m-%d %H:%M")],
                ["Rates", rate_note(trip)[0].upper() + rate_note(trip)[1:]
                          + " (each row's source is in the Rate source column)"],
                ["Record", "trips/data/expenses.csv is the record; this workbook is regenerated from it"]]:
        wi.append(row)
        wi.cell(wi.max_row, 1).font = bold
    wi.column_dimensions["A"].width, wi.column_dimensions["B"].width = 16, 90

    path = out_dir(trip) / f"{trip['name']} expenses.xlsx"
    wb.save(path)
    return path


# -------------------------------------------------------------------- commands
def send(chat: str, *args: str, text: str | None = None) -> None:
    r = subprocess.run([sys.executable, str(HERE / "send.py"), "--chat", chat, *args],
                       input=text, capture_output=True, text=True)
    if r.returncode:
        raise SystemExit(f"report: sending failed: {r.stderr.strip()}")


def post_daily(trip: dict, day: str, mark: bool, to: str | None = None) -> int:
    """Send text, chart and workbook to the trip's group (or `to`); return entries reported."""
    text, new = daily_text(trip, day)
    png, xlsx = render_chart(trip, day), build_excel(trip)
    dest = to or trip["chat_id"]
    send(dest, text=text)
    send(dest, "--photo", str(png))
    send(dest, "--document", str(xlsx), "--caption", f"{trip['name']}: all payments so far")
    if mark:
        trip["last_report"] = datetime.now(sl.SGT).isoformat(timespec="seconds")
    return new


def nightly(at: str) -> int:
    hh, mm = (int(x) for x in at.split(":"))
    trips = sl.load_trips()
    changed = False
    for trip in [t for t in trips if t["status"] == "open"]:
        if sl.phase(trip) == "planning":  # deposits only; the nightly breakdown starts when the trip does
            print(f"{trip['name']}: planning; no nightly report")
            continue
        now = datetime.now(sl.trip_tz(trip))
        today = now.date().isoformat()
        if (now.hour, now.minute) < (hh, mm) or trip.get("last_report_day") == today:
            print(f"{trip['name']}: not due ({now:%H:%M} local)")
            continue
        _, new = daily_text(trip, today)
        if new:
            post_daily(trip, today, mark=True)
            print(f"{trip['name']}: nightly report sent ({new} new entries)")
        else:
            print(f"{trip['name']}: nothing logged since the last report; nothing sent")
        trip["last_report_day"] = today
        changed = True
    if changed:
        sl.save_trips(trips)
    return 0


def end_trip(chat: str, to: str | None, send_it: bool) -> int:
    """/endtrip: close the trip, then the final settle-up, chart, workbook and Wrapped."""
    import wrapped
    trips = sl.load_trips()
    trip = sl.open_trip(chat, trips)
    settle = sl.settle_text(trip)
    trip["status"], trip["closed"] = "closed", datetime.now(sl.SGT).date().isoformat()
    sl.save_trips(trips)
    print(f"{trip['name']}: trip closed")
    dest = to or trip["chat_id"]
    out = [("text", f"Trip ended.\n{settle}")]
    if expenses(trip):
        out += [("photo", str(render_chart(trip, trip["closed"]))),
                ("document", str(build_excel(trip)), f"{trip['name']}: every payment, final"),
                ("document", str(wrapped.render(trip)), wrapped.caption(trip))]
    for kind, *rest in out:
        if not send_it:
            print(f"{kind}: {rest[0]}")
        elif kind == "text":
            send(dest, text=rest[0])
        elif kind == "photo":
            send(dest, "--photo", rest[0])
        else:
            send(dest, "--document", rest[0], "--caption", rest[1])
    return 0


def latest_trip(chat: str, trips: list[dict]) -> dict:
    """The group's open trip, or else the one it closed most recently."""
    try:
        return sl.open_trip(chat, trips)
    except sl.SplitError:
        done = [t for t in trips if t["chat_id"] == str(chat)]
        if not done:
            raise
        return max(done, key=lambda t: (t.get("closed") or "", t["created"]))


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("what", choices=["daily", "nightly", "chart", "excel", "wrapped", "endtrip", "catalogue"])
    ap.add_argument("--chat")
    ap.add_argument("--at", default="21:45", help="nightly: local time after which the report is due")
    ap.add_argument("--date", help="report day, YYYY-MM-DD (default: today on the trip's clock)")
    ap.add_argument("--send", action="store_true", help="post to the group")
    ap.add_argument("--to", help="deliver here instead of the trip's group (e.g. the owner's private chat)")
    ap.add_argument("--update-id", help="the request's Telegram update id: sends at most once per id and kind")
    ap.add_argument("--skip-empty", action="store_true",
                    help="daily: send nothing if nothing was logged since the last report")
    ap.add_argument("--no-mark", action="store_true", help="daily: don't record this as the last report")
    a = ap.parse_args(argv)
    if a.what == "nightly":
        return nightly(a.at)
    if a.send and a.update_id:
        # One answer per request: a retry, or a second copy of the same event,
        # must not post the chart or workbook again.
        marker = Path(os.environ.get("SPLITBOT_HOME", Path.home() / ".splitbot")) / "sent" / f"{a.update_id}-{a.what}"
        if marker.exists():
            print(f"already sent {a.what} for update {a.update_id}; not sending again")
            return 0
        marker.parent.mkdir(parents=True, exist_ok=True)
        marker.write_text(datetime.now(sl.SGT).isoformat(timespec="seconds"))
    if not a.chat:
        ap.error("--chat is required")
    if a.what == "endtrip":
        try:
            return end_trip(a.chat, a.to, a.send)
        except sl.SplitError as e:
            print(f"error: {e}")
            return 1
    try:
        trips = sl.load_trips()
        trip = sl.open_trip(a.chat, trips) if a.what not in ("wrapped", "catalogue") else latest_trip(a.chat, trips)
    except sl.SplitError as e:
        print(f"error: {e}")
        return 1
    day = a.date or sl.trip_today(trip)

    if a.what == "catalogue":
        import catalogue
        pdf = catalogue.render(trip)
        print(f"catalogue: {pdf}")
        if a.send:
            send(a.to or a.chat, "--document", str(pdf), "--caption", catalogue.caption(trip))
        return 0
    if a.what == "wrapped":
        import wrapped
        try:
            pdf = wrapped.render(trip)
        except sl.SplitError as e:
            print(f"error: {e}")
            return 1
        print(f"wrapped: {pdf}")
        if a.send:
            send(a.to or a.chat, "--document", str(pdf), "--caption", wrapped.caption(trip))
        return 0
    if a.what == "excel":
        path = build_excel(trip)
        print(f"excel: {path}")
        if a.send:
            send(a.to or a.chat, "--document", str(path), "--caption", f"{trip['name']}: all payments so far")
        return 0
    if a.what == "chart":
        png = render_chart(trip, day)
        print(f"chart: {png}")
        if a.send:
            send(a.to or a.chat, "--photo", str(png))
        return 0

    text, new = daily_text(trip, day)
    if a.skip_empty and new == 0:
        print("nothing new since the last report; nothing sent")
        return 0
    if a.send:
        mark = not a.no_mark and not a.to  # only a report to the group counts as the nightly one
        post_daily(trip, day, mark=mark, to=a.to)
        if mark:
            sl.save_trips(trips)
        print(text)
        return 0
    png, xlsx = render_chart(trip, day), build_excel(trip)
    print(text)
    print(f"\nchart: {png}\nexcel: {xlsx}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
